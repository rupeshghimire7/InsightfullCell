"""Phase 6 engine pieces: map outlines and region maps (checked like charts), the dashboard's
tiles, Vega-Lite specs, drill-down subsets and the HTML export."""

from __future__ import annotations

import re
from datetime import UTC, datetime

import pytest

from insightfulcell.engine.analysis import dashboard as dashboard_doc
from insightfulcell.engine.analysis import table as analysis_table
from insightfulcell.engine.analysis import vegalite
from insightfulcell.engine.analysis.charts import check_map
from insightfulcell.engine.analysis.claims import RowsClaim, ShareClaim
from insightfulcell.engine.analysis.facts import ChartData, ChartPoint, Fact, MapData
from insightfulcell.engine.analysis.miners import Finding
from insightfulcell.engine.analysis.narrate import Narration, Sentence
from insightfulcell.engine.analysis.regions import place_maps
from insightfulcell.engine.analysis.report import Preparation, Scope, build
from insightfulcell.engine.geo.maps import MAX_HEIGHT, WIDTH, outlines
from insightfulcell.engine.geo.resolve import resolve_plan
from insightfulcell.engine.planning.language import (
    MinerKind,
    ResolvedAnalysis,
    ResolvedTarget,
    Role,
    Source,
    TargetKind,
)
from insightfulcell.engine.subset import Subset, explorable, keep
from insightfulcell.engine.types import Meaning, PersonalData
from insightfulcell.exports.document import Document, build_document
from insightfulcell.exports.html import filename, render_report
from insightfulcell.exports.svg import chart_svg, format_value, map_svg
from tests.unit.engine_support import table
from tests.unit.test_engine_analysis import column, plan_with

_PATH = re.compile(r"^(M[-0-9. L]+Z)+$")
STATES = ["NSW", "New South Wales", "VIC", "Victoria", "QLD", "Tasmania"]


# --------------------------------------------------------------------------- outlines


@pytest.mark.parametrize(
    ("level", "country"), [("country", None), ("state", "AU"), ("state", "US"), ("state", "RU")]
)
def test_outlines_fit_the_view_box_as_valid_paths(level: str, country: str | None) -> None:
    drawn = outlines(level, country)
    assert drawn.width <= WIDTH + 0.1
    assert drawn.height <= MAX_HEIGHT + 0.1
    assert drawn.shapes
    for shape in drawn.shapes:
        assert _PATH.match(shape.path), shape.code
        numbers = [float(n) for n in re.findall(r"-?\d+\.\d", shape.path)]
        assert min(numbers) >= -0.1
        xs, ys = numbers[0::2], numbers[1::2]
        assert max(xs) <= drawn.width + 0.1
        assert max(ys) <= drawn.height + 0.1


def test_alaska_and_hawaii_are_insets_and_russia_stays_in_one_piece() -> None:
    us = outlines("state", "US")
    assert {"US-AK", "US-HI", "US-CA"} <= {s.code for s in us.shapes}
    assert us.width > us.height  # not stretched across the Pacific
    ru = outlines("state", "RU")
    assert ru.width / ru.height < 3  # the far east is not split to the other edge


def test_a_world_map_is_cropped_to_the_countries_with_data() -> None:
    world = outlines("country")
    europe = outlines("country", focus={"FR", "DE", "IT", "ES"})
    assert "AQ" not in {s.code for s in world.shapes}
    assert len(europe.shapes) < len(world.shapes)
    assert {"FR", "DE", "IT", "ES"} <= {s.code for s in europe.shapes}


# --------------------------------------------------------------------------- region maps


def _place_plan(target: bool = True):  # type: ignore[no-untyped-def]
    plan = plan_with(
        [
            column("state", Meaning.NOMINAL).model_copy(
                update={"meaning": Meaning.GEOGRAPHY, "role": Role.GEOGRAPHY}
            ),
            column("sales", Meaning.CURRENCY, unit="$"),
        ],
        (MinerKind.SEGMENT, ("state", "sales")),
    )
    if target:
        plan = plan.model_copy(
            update={
                "target": ResolvedTarget(
                    column="sales", kind=TargetKind.NUMERIC, rationale="", source=Source.RULES
                )
            }
        )
    return plan


def _place_rows() -> list[tuple[object, ...]]:
    rows: list[tuple[object, ...]] = []
    for i in range(80):
        state = STATES[i % 4]  # NSW and VIC, 40 rows each
        rows.append((state, 100.0 + (20 if "V" in state else 0) + i % 7))
    rows += [("Tasmania", 50.0 + i) for i in range(4)]  # a small area
    rows += [("QLD", None) for _ in range(12)]  # rows, but no sales
    return rows


def test_region_maps_show_shares_and_averages_and_hatch_small_areas() -> None:
    plan = _place_plan()
    with table([("state", "VARCHAR"), ("sales", "DOUBLE")], _place_rows()) as con:
        con.execute("CREATE VIEW cleaned AS SELECT * FROM t")
        places = resolve_plan(con, "cleaned", plan)
        derived = analysis_table.build(
            con, source="cleaned", resolutions=places, features=[], labels={"state": "State"}
        )
        maps = place_maps(con, plan, places, derived)
        assert [m.analysis_id for m in maps] == ["places:state", "places:state"]
        share, mean = (m.map for m in maps)
        for data in (share, mean):
            assert check_map(con, data).verified
        regions = {r.code: r for r in share.regions}
        assert regions["AU-NSW"].status == "value"
        assert regions["AU-NSW"].value == pytest.approx(40 / 96)
        assert regions["AU-NSW"].key == "New South Wales"
        assert regions["AU-TAS"].status == "few_rows"
        assert (regions["AU-TAS"].value, regions["AU-TAS"].n) == (None, None)
        assert regions["AU-WA"].status == "no_rows"
        assert share.drill_column == "state"
        assert share.place_label == "State"
        averages = {r.code: r for r in mean.regions}
        assert averages["AU-VIC"].value == pytest.approx(
            sum(120.0 + i % 7 for i in range(80) if "V" in STATES[i % 4]) / 40
        )
        assert averages["AU-QLD"].status == "few_rows"  # 12 rows, none with sales
        # A tampered region is caught; a hatched region must really be small.
        nsw = regions["AU-NSW"]
        bad = nsw.model_copy(update={"value": (nsw.value or 0) + 0.01})
        tampered = share.model_copy(
            update={"regions": tuple(bad if r.code == "AU-NSW" else r for r in share.regions)}
        )
        assert not check_map(con, tampered).verified
        hidden = nsw.model_copy(
            update={"status": "few_rows", "value": None, "n": None,
                    "claim": RowsClaim(column=nsw.claim.column, key="AU-NSW")}  # type: ignore[union-attr]
        )  # fmt: skip
        hiding = share.model_copy(
            update={"regions": tuple(hidden if r.code == "AU-NSW" else r for r in share.regions)}
        )
        assert not check_map(con, hiding).verified
        assert isinstance(nsw.claim, ShareClaim)


def test_without_a_numeric_outcome_there_is_one_map() -> None:
    plan = _place_plan(target=False)
    with table([("state", "VARCHAR"), ("sales", "DOUBLE")], _place_rows()) as con:
        con.execute("CREATE VIEW cleaned AS SELECT * FROM t")
        places = resolve_plan(con, "cleaned", plan)
        derived = analysis_table.build(
            con, source="cleaned", resolutions=places, features=[], labels={"state": "State"}
        )
        assert len(place_maps(con, plan, places, derived)) == 1


# --------------------------------------------------------------------------- subsets


def test_subsets_keep_the_rows_of_a_group_or_a_region() -> None:
    plan = _place_plan()
    plan = plan.model_copy(
        update={
            "dictionary": (
                *plan.dictionary,
                column("channel", Meaning.NOMINAL),
                column("email", Meaning.TEXT).model_copy(
                    update={"personal_data": PersonalData.EMAIL}
                ),
            )
        }
    )
    rows = [(s, 10.0, "Online" if i % 3 else " Store ", "a@b.c") for i, s in enumerate(STATES * 10)]
    with table(
        [("state", "VARCHAR"), ("sales", "DOUBLE"), ("channel", "VARCHAR"), ("email", "VARCHAR")],
        rows,
    ) as con:
        con.execute("CREATE TABLE cleaned AS SELECT * FROM t")

        def count(subset: Subset) -> int:
            condition = keep(con, "cleaned", plan, subset)
            return int(con.execute(f"SELECT count(*) FROM cleaned WHERE {condition}").fetchone()[0])  # type: ignore[index]

        # "NSW" and "New South Wales" are one region, as the report grouped them.
        assert count(Subset(column="state", values=("New South Wales",))) == 20
        assert count(Subset(column="state", values=("NSW",))) == 0
        assert count(Subset(column="channel", values=("Store",))) == 20  # trimmed, like groups
        assert count(Subset(column="channel", values=("Online", "Store"))) == 60
        # Values are bound, never spliced into SQL.
        assert count(Subset(column="channel", values=("x' OR '1'='1",))) == 0
    assert explorable(plan, "state")
    assert explorable(plan, "channel")
    assert not explorable(plan, "sales")
    assert not explorable(plan, "email")  # personal data
    assert not explorable(plan, "missing")
    assert Subset(column="c", values=("b", "a", "b")).fingerprint() == {
        "column": "c",
        "values": ["a", "b"],
    }


# --------------------------------------------------------------------------- dashboard


def _fact(fid: str, kind: str, numbers: tuple[str, ...], **extra: object) -> Fact:
    return Fact(
        id=fid,
        analysis_id=str(extra.pop("analysis_id", f"a:{fid}")),
        kind=kind,
        columns=("x",),
        rows=100,
        method="m",
        statement=" ".join(numbers) or "s",
        numbers=numbers,
        tier=extra.pop("tier", "strong"),
        **extra,
    )


def test_headlines_take_one_of_the_facts_own_numbers() -> None:
    group = _fact("F1", "group_comparison", ("$12", "$4", "$8", "0.4"), labels=("North", "South"))
    assert dashboard_doc.headline_of(group) == ("$12", "Highest group average: North")
    trend = _fact("F2", "trend", ("10", "15", "16", "+50%", "12", "0.3"), values={"first": 10})
    assert dashboard_doc.headline_of(trend) == ("+50%", "Change over the period")
    flat = _fact("F3", "trend", ("0", "15", "16", "12", "0.3"), values={"first": 0})
    assert dashboard_doc.headline_of(flat) == ("15", "Latest average")
    calm = _fact("F4", "outliers", ("1", "9"), values={"outside": 0})
    assert dashboard_doc.headline_of(calm) is None  # its numbers are a range, not a count
    assert dashboard_doc.headline_of(_fact("F5", "correlation", ())) is None


def _chart(title: str) -> ChartData:
    return ChartData(
        kind="bar",
        title=title,
        x_label="Region",
        y_label="Sales",
        value_format="number",
        points=(ChartPoint(label="<b>North</b>", value=12.0, n=40),),
        check={"claims": 1, "verified": True},
    )


def _report(scope: Scope | None = None):  # type: ignore[no-untyped-def]
    plan = _place_plan()
    facts = [
        _fact("F1", "group_comparison", ("$12", "$4", "$8", "0.40"), labels=("North",),
              analysis_id="seg"),
        _fact("F2", "distribution", ("$9", "$1", "$20"), tier="descriptive", analysis_id="dist"),
    ]  # fmt: skip
    analyses = [
        ResolvedAnalysis(id="seg", kind=MinerKind.SEGMENT, columns=("state", "sales"),
                         question="Does <i>sales</i> differ?", priority=1, source=Source.RULES),
        ResolvedAnalysis(id="dist", kind=MinerKind.DISTRIBUTION, columns=("sales",),
                         question="How is sales spread?", priority=2, source=Source.RULES),
    ]  # fmt: skip
    findings = [
        Finding(analysis=analyses[0], facts=[facts[0]], chart=_chart("Sales <script>")),
        Finding(analysis=analyses[1], facts=[facts[1]], chart=_chart("Spread")),
    ]
    narration = Narration(
        summary=(Sentence(text="North sells most.", citations=("F1",)),),
        sections={"seg": (Sentence(text="North leads <img src=x>.", citations=("F1",)),),
                  "dist": (Sentence(text="Sales vary.", citations=("F2",)),)},
        written_by={"seg": "model", "dist": "engine"},
        source="mixed",
    )  # fmt: skip
    report = build(
        dataset_name='sales "2026" <x>.csv',
        plan=plan,
        findings=findings,
        narration=narration,
        preparation=Preparation(
            rows_before=10, rows_after=10, applied=(), not_applied=0, reviewed_by_person=True
        ),
        rows=10,
        highlights=["F1", "F2"],
        scope=scope,
    )
    return report, {f.id: f for f in facts}


def test_the_dashboard_puts_findings_first_and_descriptions_in_the_gallery() -> None:
    report, facts = _report()
    board = dashboard_doc.build(report, facts, charts_withheld=1)
    assert [t.fact_id for t in board.tiles] == ["F1", "F2"]
    assert [i.analysis_id for i in board.insights] == ["seg"]
    assert [i.analysis_id for i in board.gallery] == ["dist"]
    assert [q.analysis_id for q in board.questions] == ["seg", "dist"]
    assert board.questions[0].answer == report.sections[0].answer
    assert board.glance.charts_withheld == 1
    assert board.exploratory is False
    # An answer whose test did not survive the false-discovery check is not a top finding.
    unreliable = {"F1": facts["F1"].model_copy(update={"tier": "inconclusive"}), "F2": facts["F2"]}
    weak = dashboard_doc.build(report, unreliable)
    assert weak.insights == ()
    assert [i.analysis_id for i in weak.gallery] == ["seg", "dist"]
    scoped, _ = _report(Scope(parent_run_id="p", column="state", display_name="State",
                              values=("Victoria",), rows=40))  # fmt: skip
    assert scoped.title.startswith("Exploring State: Victoria")
    explored = dashboard_doc.build(scoped, facts)
    assert explored.exploratory is True
    assert explored.title.startswith("Dashboard: Exploring State: Victoria")
    assert board.title == 'Dashboard: sales "2026" <x>.csv'


# --------------------------------------------------------------------------- exports


def test_vega_lite_specs_carry_the_values_inline() -> None:
    spec = vegalite.spec(_chart("Sales"))
    assert spec["$schema"].startswith("https://vega.github.io/schema/vega-lite/v5")
    assert spec["data"]["values"] == [{"label": "<b>North</b>", "value": 12.0, "n": 40}]
    assert spec["mark"]["type"] == "bar"


def test_the_html_export_escapes_everything_and_runs_nothing() -> None:
    report, facts = _report()
    board = dashboard_doc.build(report, facts)
    fact_rows = [
        {**f.model_dump(mode="json"), "columns": [{"name": "x", "display_name": "X <u>"}],
         "test": None, "check": None, "content_hash": "0" * 64, "exploratory": False}
        for f in facts.values()
    ]  # fmt: skip
    html = render_report(
        report=report,
        dashboard=board,
        facts=fact_rows,
        run_id="0192f3a4-0000-7000-8000-000000000000",
        written_at=datetime(2026, 10, 11, 9, 30, tzinfo=UTC),
    )
    assert "<script" not in html
    assert "<img" not in html
    assert "<b>North</b>" not in html
    assert "&lt;b&gt;North&lt;/b&gt;" in html
    assert "&lt;img src=x&gt;" in html
    assert "Does &lt;i&gt;sales&lt;/i&gt; differ?" in html
    assert "X &lt;u&gt;" in html
    assert 'id="fact-F1"' in html
    assert 'href="#fact-F1"' in html
    assert "11 October 2026" in html


def _document() -> Document:
    report, facts = _report()
    board = dashboard_doc.build(report, facts)
    fact_rows = [
        {**f.model_dump(mode="json"), "columns": [{"name": "x", "display_name": "X <u>"}],
         "test": {"method": "Welch's t-test", "effect_name": "eta", "effect": 0.42,
                  "ci_low": 0.1, "ci_high": 0.6, "p_value": 0.001, "q_value": 0.004},
         "check": None, "content_hash": "0" * 64, "exploratory": False}
        for f in facts.values()
    ]  # fmt: skip
    return build_document(
        report=report,
        dashboard=board,
        facts=fact_rows,
        run_id="0192f3a4-0000-7000-8000-000000000000",
        written_at=datetime(2026, 10, 11, 9, 30, tzinfo=UTC),
    )


def test_the_markdown_export_escapes_everything() -> None:
    from insightfulcell.exports.markdown import escape, render_markdown  # noqa: PLC0415

    text = render_markdown(_document())
    for tag in ("<b>", "<img", "<i>"):
        assert tag not in text
    assert "&lt;b&gt;North&lt;/b&gt;" in text
    assert "Does &lt;i&gt;sales&lt;/i&gt; differ?" in text
    assert "\\[F1\\]" in text  # source tags, escaped so they never become links
    assert ":--- |" in text  # the numbers behind each chart
    assert "data:image/svg+xml;base64," in text
    assert "Content hash 0000" in text
    assert "11 October 2026" in text
    # A value can neither start a block nor format text, and stays on one line.
    assert escape("# 1. *bold* [x](y)\nnext") == "\\# 1. \\*bold\\* \\[x\\](y) next"


def test_the_word_export_is_a_document_with_the_report_in_it() -> None:
    import io  # noqa: PLC0415

    import docx  # noqa: PLC0415

    from insightfulcell.exports.word import clean, render_docx  # noqa: PLC0415

    content = render_docx(_document())
    assert content[:2] == b"PK"  # a zip package
    opened = docx.Document(io.BytesIO(content))
    text = "\n".join(p.text for p in opened.paragraphs)
    assert "Does <i>sales</i> differ?" in text  # text, never markup
    assert "Facts" in text
    assert len(opened.inline_shapes) >= 1  # the chart, as a picture
    cells = [c.text for t in opened.tables for row in t.rows for c in row.cells]
    assert "<b>North</b>" in cells
    assert opened.core_properties.title == _document().title
    assert clean("a\x00b\x0bc") == "a b c"


def test_the_pdf_export_lays_out_pages() -> None:
    from insightfulcell.exports.pdf import render_pdf  # noqa: PLC0415

    content = render_pdf(_document())
    assert content.startswith(b"%PDF-")
    assert content.rstrip().endswith(b"%%EOF")
    assert b"/Subtype /Image" in content  # the chart
    assert b"DejaVu" in content  # the bundled font, so every character prints


def test_export_file_names_and_values() -> None:
    assert (
        filename('Sales "Q3" <x>/../é.csv', "0192f3a4-aaaa")
        == "report-sales-q3-x-e-csv-0192f3a4.html"
    )
    assert filename("", "abcdef12-3") == "report-dataset-abcdef12.html"
    assert filename("Sales", "abcdef12-3", "pdf") == "report-sales-abcdef12.pdf"
    assert format_value(0.4167, "percent", None) == "42%"
    assert format_value(-0.05, "signed_percent", None) == "-5.0%"
    assert format_value(1234.5, "number", "$") == "$1,234"
    assert format_value(float("nan"), "number", None) == "n/a"
    assert "<title>" in chart_svg(_chart("A & B"))
    assert "A &amp; B" in chart_svg(_chart("A & B"))
    empty = MapData(title="M", value_label="v", value_format="percent", level="state",
                    width=10, height=10, regions=())  # fmt: skip
    assert map_svg(empty).startswith("<svg")
