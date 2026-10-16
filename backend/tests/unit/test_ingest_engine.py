from __future__ import annotations

import hashlib
import json
import zipfile
from datetime import date
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook

from insightfulcell.domain.errors import (
    ArchiveBomb,
    ContentMismatch,
    FileTooLarge,
    IntegrityMismatch,
    ParseFailure,
    ParseTimeout,
    TooManyColumns,
    TooManyRows,
)
from insightfulcell.ingestion import engine
from insightfulcell.ingestion.engine import ingest_file, sha256_file
from insightfulcell.ingestion.formats import FileFormat
from insightfulcell.ingestion.limits import IngestLimits


def _types(result_columns: tuple[object, ...]) -> dict[str, str]:
    return {c.name: c.type for c in result_columns}  # type: ignore[attr-defined]


def _read_parquet_rows(path: Path) -> list[tuple[object, ...]]:
    return duckdb.sql(f"SELECT * FROM read_parquet('{path}')").fetchall()


class TestCsv:
    def test_infers_types_and_writes_parquet(self, tmp_path: Path) -> None:
        src = tmp_path / "sales.csv"
        src.write_text(
            "id,price,sold_on,suburb\n1,10.5,2026-01-02,Newtown\n2,20,2026-02-03,Glebe\n"
        )
        result = ingest_file(src, FileFormat.CSV, tmp_path / "out" / "sales.parquet")

        assert result.rows == 2
        types = _types(result.columns)
        assert types["id"] == "BIGINT"
        assert types["price"] == "DOUBLE"
        assert types["sold_on"] == "DATE"
        assert types["suburb"] == "VARCHAR"
        assert result.parquet_path.exists()
        assert result.source_sha256 == hashlib.sha256(src.read_bytes()).hexdigest()
        assert not result.warnings
        assert len(_read_parquet_rows(result.parquet_path)) == 2

    def test_no_partial_or_scratch_files_left_behind(self, tmp_path: Path) -> None:
        src = tmp_path / "a.csv"
        src.write_text("a,b\n1,2\n")
        out = tmp_path / "out"
        ingest_file(src, FileFormat.CSV, out / "a.parquet")
        assert sorted(p.name for p in out.iterdir()) == ["a.parquet"]

    def test_falls_back_to_text_when_late_rows_break_inference(self, tmp_path: Path) -> None:
        src = tmp_path / "late.csv"
        rows = ["n"] + ["1"] * 150_000 + ["not-a-number"]
        src.write_text("\n".join(rows) + "\n")
        result = ingest_file(src, FileFormat.CSV, tmp_path / "late.parquet")

        assert _types(result.columns)["n"] == "VARCHAR"
        assert result.rows == 150_001
        assert any("read as text" in w for w in result.warnings)

    def test_latin1_encoding_is_detected(self, tmp_path: Path) -> None:
        src = tmp_path / "latin.csv"
        src.write_bytes("name,city\nJos\xe9,Z\xfcrich\n".encode("latin-1"))
        result = ingest_file(src, FileFormat.CSV, tmp_path / "latin.parquet")
        assert _read_parquet_rows(result.parquet_path) == [("José", "Zürich")]

    def test_utf16_encoding_is_detected(self, tmp_path: Path) -> None:
        src = tmp_path / "u16.csv"
        src.write_bytes("a,b\n1,x\n".encode("utf-16"))
        result = ingest_file(src, FileFormat.CSV, tmp_path / "u16.parquet")
        assert result.rows == 1
        assert [c.name for c in result.columns] == ["a", "b"]

    @pytest.mark.parametrize("content", ["", "a,b\n"], ids=["empty", "header-only"])
    def test_files_without_data_rows_are_rejected(self, tmp_path: Path, content: str) -> None:
        src = tmp_path / "nodata.csv"
        src.write_text(content)
        with pytest.raises(ParseFailure) as info:
            ingest_file(src, FileFormat.CSV, tmp_path / "nodata.parquet")
        assert info.value.hint
        assert not (tmp_path / "nodata.parquet").exists()


class TestMalformedDelimited:
    """One bad row used to make the reader collapse the whole file into one column, silently."""

    @staticmethod
    def _rows(bad_line: str, *, at: int = 150) -> str:
        rows = ["id,city,amount"] + [f"{i},City{i % 4},{i * 2.5}" for i in range(200)]
        rows[at] = bad_line
        return "\n".join(rows) + "\n"

    def test_a_row_with_an_extra_field_is_refused_with_its_line(self, tmp_path: Path) -> None:
        src = tmp_path / "extra.csv"
        src.write_text(self._rows("149,Secret Value,375.0,EXTRA"))
        with pytest.raises(ParseFailure) as info:
            ingest_file(src, FileFormat.CSV, tmp_path / "extra.parquet")
        assert info.value.detail == "Line 151 has 4 values, but the header has 3 columns."
        assert "separator" in (info.value.hint or "")
        assert "Secret" not in str(info.value.detail)  # never a cell value
        assert not (tmp_path / "extra.parquet").exists()

    def test_short_rows_are_counted(self, tmp_path: Path) -> None:
        src = tmp_path / "short.csv"
        src.write_text("a,b,c\n1,2,3\n4,5\n6,7,8\n9\n")
        with pytest.raises(ParseFailure) as info:
            ingest_file(src, FileFormat.CSV, tmp_path / "short.parquet")
        assert info.value.detail == (
            "Line 3 has 2 values, but the header has 3 columns. 1 other line does not match either."
        )

    def test_an_unclosed_quote_is_refused(self, tmp_path: Path) -> None:
        src = tmp_path / "quote.csv"
        src.write_text(self._rows('149,"Darwin, NT,375.0'))
        with pytest.raises(ParseFailure) as info:
            ingest_file(src, FileFormat.CSV, tmp_path / "quote.parquet")
        assert "line 151 is never closed" in str(info.value.detail)

    def test_tsv_rows_are_checked_too(self, tmp_path: Path) -> None:
        src = tmp_path / "bad.tsv"
        src.write_text("a\tb\n1\tx\n2\ty\tz\n")
        with pytest.raises(ParseFailure) as info:
            ingest_file(src, FileFormat.TSV, tmp_path / "bad.parquet")
        assert "Line 3 has 3 values" in str(info.value.detail)

    def test_well_formed_files_still_read(self, tmp_path: Path) -> None:
        src = tmp_path / "ok.csv"
        src.write_text(
            'name,note,n\n"Smith, J","said ""hi""\nthen left",1\n\nLee,plain,2\n'  # quoted
        )
        result = ingest_file(src, FileFormat.CSV, tmp_path / "ok.parquet")
        assert result.rows == 2
        assert [c.name for c in result.columns] == ["name", "note", "n"]

    def test_the_header_decides_the_separator(self, tmp_path: Path) -> None:
        src = tmp_path / "semi.csv"
        src.write_text('"Price, AUD";Qty\n"1,5";2\n"2,0";3\n')
        result = ingest_file(src, FileFormat.CSV, tmp_path / "semi.parquet")
        assert [c.name for c in result.columns] == ["Price, AUD", "Qty"]

    def test_one_column_files_need_no_separator(self, tmp_path: Path) -> None:
        src = tmp_path / "one.csv"
        src.write_text("value\n1\n2\n3\n")
        assert ingest_file(src, FileFormat.CSV, tmp_path / "one.parquet").rows == 3

    def test_a_table_that_still_collapses_is_refused(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        src = tmp_path / "split.csv"
        src.write_text("a,b\n1,2\n3,4\n")
        real = engine._open_relation

        def one_column(con, source, fmt, *, all_varchar):  # type: ignore[no-untyped-def]
            del fmt
            return real(con, source, FileFormat.TSV, all_varchar=all_varchar)

        monkeypatch.setattr(engine, "_open_relation", one_column)
        with pytest.raises(ParseFailure) as info:
            ingest_file(src, FileFormat.CSV, tmp_path / "split.parquet")
        assert "could not be split" in str(info.value.detail)


class TestTsvJsonParquet:
    def test_tsv(self, tmp_path: Path) -> None:
        src = tmp_path / "t.tsv"
        src.write_text("a\tb\n1\tx\n2\ty\n")
        result = ingest_file(src, FileFormat.TSV, tmp_path / "t.parquet")
        assert result.rows == 2
        assert [c.name for c in result.columns] == ["a", "b"]

    def test_json_array_and_ndjson(self, tmp_path: Path) -> None:
        arr = tmp_path / "arr.json"
        arr.write_text(json.dumps([{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]))
        nd = tmp_path / "nd.json"
        nd.write_text('{"a":1,"b":"x"}\n{"a":2,"b":"y"}\n')
        for src in (arr, nd):
            result = ingest_file(src, FileFormat.JSON, tmp_path / f"{src.stem}.parquet")
            assert result.rows == 2
            assert {c.name for c in result.columns} == {"a", "b"}

    def test_parquet_roundtrip(self, tmp_path: Path) -> None:
        src = tmp_path / "in.parquet"
        duckdb.sql("SELECT 1 AS a, 'x' AS b UNION ALL SELECT 2, 'y'").write_parquet(str(src))
        result = ingest_file(src, FileFormat.PARQUET, tmp_path / "out.parquet")
        assert result.rows == 2

    def test_invalid_json_fails_without_echoing_cell_values(self, tmp_path: Path) -> None:
        src = tmp_path / "broken.json"
        src.write_text('[{"email": "secret@example.com", "n": 1}, {"email": "oops"')
        with pytest.raises(ParseFailure) as info:
            ingest_file(src, FileFormat.JSON, tmp_path / "broken.parquet")
        assert "secret@example.com" not in info.value.detail
        assert "secret@example.com" not in (info.value.hint or "")
        assert info.value.hint


class TestExcel:
    def test_first_sheet_is_streamed_with_types(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["id", "amount", "day", "flag"])
        ws.append([1, 9.5, date(2026, 1, 2), True])
        ws.append([2, 10.0, date(2026, 1, 3), False])
        second = wb.create_sheet("ignored")
        second.append(["z"])
        src = tmp_path / "book.xlsx"
        wb.save(src)

        result = ingest_file(src, FileFormat.XLSX, tmp_path / "book.parquet")
        assert result.rows == 2
        assert [c.name for c in result.columns] == ["id", "amount", "day", "flag"]
        # Excel stores dates as datetimes; the profile stage can later narrow all-midnight values.
        assert _types(result.columns)["day"] == "TIMESTAMP"

    def test_blank_header_cells_get_names_and_trailing_blank_columns_are_dropped(
        self, tmp_path: Path
    ) -> None:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["a", None, "c", None, None])
        ws.append([1, 2, 3, None, None])
        src = tmp_path / "b.xlsx"
        wb.save(src)
        result = ingest_file(src, FileFormat.XLSX, tmp_path / "b.parquet")
        assert [c.name for c in result.columns] == ["a", "column1", "c"]

    def test_too_many_rows(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["a"])
        for i in range(20):
            ws.append([i])
        src = tmp_path / "r.xlsx"
        wb.save(src)
        with pytest.raises(TooManyRows):
            ingest_file(
                src, FileFormat.XLSX, tmp_path / "r.parquet", limits=IngestLimits(max_rows=10)
            )

    def test_empty_sheet_is_rejected(self, tmp_path: Path) -> None:
        src = tmp_path / "e.xlsx"
        Workbook().save(src)
        with pytest.raises(ParseFailure):
            ingest_file(src, FileFormat.XLSX, tmp_path / "e.parquet")

    def test_zip_bomb_is_rejected_before_parsing(self, tmp_path: Path) -> None:
        src = tmp_path / "bomb.xlsx"
        with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", b"0" * 5_000_000)
        with pytest.raises(ArchiveBomb):
            ingest_file(src, FileFormat.XLSX, tmp_path / "bomb.parquet")

    def test_too_many_entries_is_rejected(self, tmp_path: Path) -> None:
        src = tmp_path / "many.xlsx"
        with zipfile.ZipFile(src, "w") as archive:
            for i in range(30):
                archive.writestr(f"f{i}.txt", b"x")
        with pytest.raises(ArchiveBomb):
            ingest_file(
                src,
                FileFormat.XLSX,
                tmp_path / "many.parquet",
                limits=IngestLimits(max_zip_entries=10),
            )

    def test_corrupt_zip_is_a_content_mismatch(self, tmp_path: Path) -> None:
        src = tmp_path / "corrupt.xlsx"
        src.write_bytes(b"PK\x03\x04" + b"garbage" * 50)
        with pytest.raises(ContentMismatch):
            ingest_file(src, FileFormat.XLSX, tmp_path / "corrupt.parquet")


class TestLimitsAndIntegrity:
    def test_file_size_limit_is_checked_first(self, tmp_path: Path) -> None:
        src = tmp_path / "big.csv"
        src.write_text("a\n" + "1\n" * 100)
        with pytest.raises(FileTooLarge):
            ingest_file(
                src, FileFormat.CSV, tmp_path / "o.parquet", limits=IngestLimits(max_file_bytes=10)
            )

    def test_row_cap(self, tmp_path: Path) -> None:
        src = tmp_path / "rows.csv"
        src.write_text("a\n" + "1\n" * 50)
        with pytest.raises(TooManyRows):
            ingest_file(
                src, FileFormat.CSV, tmp_path / "o.parquet", limits=IngestLimits(max_rows=10)
            )
        assert not (tmp_path / "o.parquet").exists()

    def test_row_cap_allows_exactly_the_limit(self, tmp_path: Path) -> None:
        src = tmp_path / "rows.csv"
        src.write_text("a\n" + "1\n" * 10)
        result = ingest_file(
            src, FileFormat.CSV, tmp_path / "o.parquet", limits=IngestLimits(max_rows=10)
        )
        assert result.rows == 10

    def test_column_cap(self, tmp_path: Path) -> None:
        src = tmp_path / "wide.csv"
        src.write_text(",".join(f"c{i}" for i in range(30)) + "\n" + ",".join("1" * 30) + "\n")
        with pytest.raises(TooManyColumns):
            ingest_file(
                src, FileFormat.CSV, tmp_path / "o.parquet", limits=IngestLimits(max_columns=10)
            )

    def test_long_column_name_is_rejected(self, tmp_path: Path) -> None:
        src = tmp_path / "name.csv"
        src.write_text("x" * 300 + ",b\n1,2\n")
        with pytest.raises(ParseFailure):
            ingest_file(src, FileFormat.CSV, tmp_path / "o.parquet")

    def test_checksum_mismatch_is_rejected(self, tmp_path: Path) -> None:
        src = tmp_path / "a.csv"
        src.write_text("a\n1\n")
        with pytest.raises(IntegrityMismatch):
            ingest_file(src, FileFormat.CSV, tmp_path / "o.parquet", expected_sha256="0" * 64)

    def test_checksum_match_is_accepted_case_insensitively(self, tmp_path: Path) -> None:
        src = tmp_path / "a.csv"
        src.write_text("a\n1\n")
        digest = sha256_file(src).upper()
        assert (
            ingest_file(src, FileFormat.CSV, tmp_path / "o.parquet", expected_sha256=digest).rows
            == 1
        )

    def test_binary_content_in_csv_is_rejected(self, tmp_path: Path) -> None:
        src = tmp_path / "fake.csv"
        src.write_bytes(b"%PDF-1.7\n" + b"\x00" * 100)
        with pytest.raises(ContentMismatch):
            ingest_file(src, FileFormat.CSV, tmp_path / "o.parquet")

    def test_parse_timeout_interrupts_a_long_running_query(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def slow(con: duckdb.DuckDBPyConnection, *_args: object, **_kwargs: object) -> None:
            con.execute("SELECT count(*) FROM range(100000000000) a, range(1000) b").fetchall()

        monkeypatch.setattr(engine, "_read_and_write", slow)
        src = tmp_path / "slow.csv"
        src.write_text("a\n1\n")
        with pytest.raises(ParseTimeout):
            ingest_file(
                src,
                FileFormat.CSV,
                tmp_path / "o.parquet",
                limits=IngestLimits(parse_timeout_seconds=0.3),
            )
