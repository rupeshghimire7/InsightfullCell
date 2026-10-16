"""The report as a Word document (.docx), built with python-docx.

Headings use Word's own heading styles (so the navigation pane and a table of contents work),
charts and maps are pictures of the same SVG the HTML export shows, each followed by the table
of its numbers. Text from the data is written as text runs, never parsed.
"""

from __future__ import annotations

import io
import re
from typing import Final

import docx
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from insightfulcell.exports.document import (
    Block,
    Bullets,
    Document,
    Figure,
    Heading,
    Note,
    Para,
    Table,
    Text,
)
from insightfulcell.exports.images import svg_to_png

ACCENT: Final = RGBColor(0x2F, 0x6F, 0x62)
MUTED: Final = RGBColor(0x5E, 0x62, 0x5E)
#: XML 1.0 refuses these; a stray control character in a value must not break the file.
_CONTROL: Final = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def clean(text: str) -> str:
    return _CONTROL.sub(" ", text)


def render_docx(doc: Document) -> bytes:
    out = docx.Document()
    style = out.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    out.core_properties.title = clean(doc.title)
    out.core_properties.subject = clean(doc.dataset_name)
    out.core_properties.author = "InsightfulCell"
    out.core_properties.comments = f"Run {doc.run_id}"

    out.add_heading(clean(doc.title), level=0)
    meta = out.add_paragraph()
    _run(meta, Text(doc.meta), small=True)
    for block in doc.blocks:
        _block(out, block)
    footer = out.sections[0].footer.paragraphs[0]
    _run(footer, Text(doc.footer), small=True)

    buffer = io.BytesIO()
    out.save(buffer)
    return buffer.getvalue()


def _run(paragraph: Paragraph, part: Text, *, small: bool = False) -> None:
    run = paragraph.add_run(clean(part.text))
    if part.cite:
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = ACCENT
        run.font.name = "Consolas"
        paragraph.add_run(" ")
        return
    run.bold = part.strong or None
    if small:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def _block(out: DocxDocument, block: Block) -> None:
    match block:
        case Heading(level=level, text=text):
            out.add_heading(clean(text), level=min(level, 3))
        case Para(parts=parts, small=small):
            paragraph = out.add_paragraph()
            for part in parts:
                _run(paragraph, part, small=small)
        case Note(title=title, text=text):
            paragraph = out.add_paragraph(style="Intense Quote")
            _run(paragraph, Text(title + " ", strong=True))
            _run(paragraph, Text(text))
        case Bullets(items=items, small=small):
            for item in items:
                paragraph = out.add_paragraph(style="List Bullet")
                for part in item:
                    _run(paragraph, part, small=small)
        case Table():
            _table(out, block)
        case Figure():
            _figure(out, block)


def _table(out: DocxDocument, table: Table) -> None:
    grid = out.add_table(rows=1, cols=len(table.header))
    grid.style = "Light Grid Accent 1"
    grid.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, name in enumerate(table.header):
        cell = grid.rows[0].cells[i]
        cell.text = clean(name)
        if table.numeric[i]:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in table.rows:
        cells = grid.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = clean(value)
            if table.numeric[i]:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    out.add_paragraph()


def _figure(out: DocxDocument, figure: Figure) -> None:
    caption = out.add_paragraph()
    _run(caption, Text(figure.title, strong=True))
    out.add_picture(io.BytesIO(svg_to_png(figure.svg)), width=Inches(6.0))
    if figure.legend:
        _run(out.add_paragraph(), Text(" · ".join(figure.legend)), small=True)
    if figure.note:
        _run(out.add_paragraph(), Text(figure.note), small=True)
    _table(out, figure.table)
