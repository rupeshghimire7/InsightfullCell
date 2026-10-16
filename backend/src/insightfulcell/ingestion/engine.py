"""Ingestion engine (IC-102, FR-INGEST-01): any supported file -> typed Parquet + a schema.

Design notes
- DuckDB does the parsing (fast, streaming, memory-limited); Excel is streamed through
  openpyxl in read-only mode into a temporary CSV so every format shares one typing path.
- Limits are enforced *before* the expensive work where possible (file size, column count) and
  *during* it otherwise (row cap via ``LIMIT max_rows + 1``), so a hostile file cannot make us
  materialise more than the limits allow.
- Parsing runs under a watchdog that interrupts DuckDB after ``parse_timeout_seconds``.
- Error messages shown to users never echo cell values: DuckDB conversion errors quote the
  offending value, which may be personal data. The original error stays on ``__cause__``.
- Nothing here knows about HTTP, databases or S3. The worker (next sprint) calls
  ``ingest_file`` with a local path.
"""

from __future__ import annotations

import csv
import hashlib
import shutil
import tempfile
import threading
from collections.abc import Iterator
from contextlib import contextmanager, suppress
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import TYPE_CHECKING, Any, Final

import duckdb

from insightfulcell.domain.errors import (
    FileTooLarge,
    IntegrityMismatch,
    ParseFailure,
    ParseTimeout,
    TooManyColumns,
    TooManyRows,
)
from insightfulcell.ingestion.formats import SNIFF_BYTES, FileFormat, sniff
from insightfulcell.ingestion.limits import IngestLimits, check_zip_archive

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

_HASH_CHUNK: Final = 1024 * 1024
_ENCODING_PROBE_BYTES: Final = 64 * 1024
_CSV_SAMPLE_ROWS: Final = 100_000
#: Separators a delimited file may use; the header decides which one (see ``_delimiter``).
_DELIMITERS: Final = (",", ";", "\t", "|")
#: Records read to choose the separator.
_DELIMITER_SAMPLE_RECORDS: Final = 200
#: A quoted value may span lines; past this size it is an unclosed quote, not a value.
_MAX_FIELD_CHARS: Final = 16 * 1024 * 1024
# Process-wide (the csv module has no per-reader limit); set once rather than per call, so
# concurrent ingests never see each other's setting.
csv.field_size_limit(max(csv.field_size_limit(), _MAX_FIELD_CHARS))


@dataclass(frozen=True, slots=True)
class ColumnSchema:
    name: str
    type: str  # DuckDB logical type, e.g. BIGINT, DOUBLE, VARCHAR, TIMESTAMP


@dataclass(slots=True)
class _ParseJob:
    """Everything one parse needs; ``warnings`` is appended to as the parse proceeds."""

    source: Path
    fmt: FileFormat
    dest: Path
    limits: IngestLimits
    warnings: list[str]


@dataclass(frozen=True, slots=True)
class IngestResult:
    format: FileFormat
    rows: int
    columns: tuple[ColumnSchema, ...]
    source_sha256: str
    source_bytes: int
    parquet_path: Path
    parquet_bytes: int
    warnings: tuple[str, ...] = field(default=())


def sha256_file(path: Path) -> str:
    """Streaming SHA-256 (hex) of a file."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(_HASH_CHUNK):
            digest.update(chunk)
    return digest.hexdigest()


def ingest_file(
    source: Path,
    fmt: FileFormat,
    dest: Path,
    *,
    limits: IngestLimits | None = None,
    expected_sha256: str | None = None,
) -> IngestResult:
    """Parse ``source`` into typed Parquet at ``dest`` and describe it.

    ``expected_sha256`` is the hash recorded when the upload completed: passing it makes this
    function the integrity check the API deliberately does not perform (ADR-026).
    """
    limits = limits or IngestLimits()
    size = source.stat().st_size
    if size > limits.max_file_bytes:
        raise FileTooLarge(
            f"The file is {size / 1024 / 1024:,.0f} MB; the limit is "
            f"{limits.max_file_bytes // 1024 // 1024} MB.",
            hint="Try Parquet for large data, or split the file.",
        )
    digest = sha256_file(source)
    if expected_sha256 is not None and digest != expected_sha256.lower():
        raise IntegrityMismatch(
            "The stored file does not match the checksum recorded at upload.",
            hint="Upload the file again.",
        )
    with source.open("rb") as handle:
        sniff(handle.read(SNIFF_BYTES), fmt)

    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp_dest = dest.with_name(dest.name + ".tmp")
    warnings: list[str] = []
    try:
        with _workdir(dest) as workdir:
            parse_source, parse_fmt = source, fmt
            header_fields = 0
            if fmt in {FileFormat.CSV, FileFormat.TSV}:
                encoding = _detect_encoding(source)
                if encoding != "utf-8":
                    # DuckDB's core build reads UTF-8 only (other encodings need a downloadable
                    # extension, which workers must not do), so transcode with bounded memory.
                    parse_source = workdir / "utf8.txt"
                    _transcode_to_utf8(source, parse_source, encoding)
                    warnings.append(
                        f"The file was not UTF-8; it was read as {encoding.upper()} text."
                    )
                header_fields = _check_structure(parse_source, fmt, limits)
            if fmt is FileFormat.XLSX:
                check_zip_archive(source, limits)
                parse_source = workdir / "sheet.csv"
                _xlsx_to_csv(source, parse_source, limits)
                parse_fmt = FileFormat.CSV
            job = _ParseJob(parse_source, parse_fmt, tmp_dest, limits, warnings)
            rows, columns = _parse_to_parquet(job)
            if fmt in {FileFormat.CSV, FileFormat.TSV} and header_fields > 1 and len(columns) == 1:
                # The rows are consistent, yet the reader did not split them: never analyse
                # a table that silently became one column.
                raise ParseFailure(
                    f"The file's header has {header_fields} columns, but it could not be split "
                    "into them.",
                    hint="Save it again as a standard CSV (comma separated, double quotes) and "
                    "upload it again.",
                )
        tmp_dest.replace(dest)
    finally:
        with suppress(FileNotFoundError):
            tmp_dest.unlink()

    return IngestResult(
        format=fmt,
        rows=rows,
        columns=columns,
        source_sha256=digest,
        source_bytes=size,
        parquet_path=dest,
        parquet_bytes=dest.stat().st_size,
        warnings=tuple(warnings),
    )


# --------------------------------------------------------------------------- parsing


def _parse_to_parquet(job: _ParseJob) -> tuple[int, tuple[ColumnSchema, ...]]:
    limits = job.limits
    con = duckdb.connect(
        ":memory:",
        config={"memory_limit": limits.duckdb_memory_limit, "threads": limits.duckdb_threads},
    )
    timed_out = threading.Event()

    def _on_timeout() -> None:
        timed_out.set()
        con.interrupt()

    watchdog = threading.Timer(limits.parse_timeout_seconds, _on_timeout)
    watchdog.daemon = True
    watchdog.start()
    try:
        try:
            return _read_and_write(con, job)
        except (TooManyColumns, TooManyRows, ParseFailure):
            raise
        except duckdb.Error as exc:
            if timed_out.is_set():
                raise ParseTimeout(
                    f"Parsing did not finish within {limits.parse_timeout_seconds:.0f} seconds.",
                    hint="Try a smaller file, or export it as Parquet.",
                ) from exc
            raise _parse_failure(job.fmt) from exc
    finally:
        watchdog.cancel()
        con.close()


def _read_and_write(
    con: duckdb.DuckDBPyConnection, job: _ParseJob
) -> tuple[int, tuple[ColumnSchema, ...]]:
    """Read ``source`` and write Parquet, retrying CSV/TSV once as all-text on a typing failure.

    Type inference samples the first rows; a later row may not fit the inferred type. Reading
    everything as text beats failing the upload, and the warning tells the profile stage why.
    """
    source, fmt, dest, limits = job.source, job.fmt, job.dest, job.limits
    is_delimited = fmt in {FileFormat.CSV, FileFormat.TSV}
    for all_varchar in (False, True) if is_delimited else (False,):
        try:
            relation = _open_relation(con, source, fmt, all_varchar=all_varchar)
            _check_columns(relation.columns, limits)
            # LIMIT max_rows + 1: one extra row proves the cap was exceeded without a full scan.
            relation.limit(limits.max_rows + 1).write_parquet(str(dest))
        except duckdb.InterruptException:
            raise
        except duckdb.Error:
            if all_varchar or not is_delimited:
                raise
            continue
        if all_varchar:
            job.warnings.append(
                "Column types could not be inferred reliably; all columns were read as text."
            )
        return _finish(con, relation, dest, limits)
    raise AssertionError("unreachable: the loop returns or raises")  # pragma: no cover


def _finish(
    con: duckdb.DuckDBPyConnection,
    relation: duckdb.DuckDBPyRelation,
    dest: Path,
    limits: IngestLimits,
) -> tuple[int, tuple[ColumnSchema, ...]]:
    count_row = con.execute("SELECT count(*) FROM read_parquet(?)", [str(dest)]).fetchone()
    rows = int(count_row[0]) if count_row else 0
    if rows == 0:
        raise ParseFailure(
            "The file has a header but no data rows.",
            hint="Check you exported the data and not only the column names.",
        )
    if rows > limits.max_rows:
        raise TooManyRows(
            f"The file has more than {limits.max_rows:,} rows.",
            hint="Upload a smaller file or an aggregated extract.",
        )
    columns = tuple(
        ColumnSchema(name=name, type=str(dtype))
        for name, dtype in zip(relation.columns, relation.types, strict=True)
    )
    return rows, columns


def _open_relation(
    con: duckdb.DuckDBPyConnection, source: Path, fmt: FileFormat, *, all_varchar: bool
) -> duckdb.DuckDBPyRelation:
    path = str(source)
    if fmt is FileFormat.PARQUET:
        return con.read_parquet(path)
    if fmt is FileFormat.JSON:
        return con.read_json(path)
    options: dict[str, Any] = {"sample_size": _CSV_SAMPLE_ROWS, "all_varchar": all_varchar}
    if fmt is FileFormat.TSV:
        options["delimiter"] = "\t"
    return con.read_csv(path, **options)


def _check_columns(names: list[str], limits: IngestLimits) -> None:
    if len(names) > limits.max_columns:
        raise TooManyColumns(
            f"The file has {len(names):,} columns; the limit is {limits.max_columns}.",
            hint="Remove columns you do not need and upload again.",
        )
    for name in names:
        if len(name) > limits.max_column_name_chars:
            raise ParseFailure(
                f"A column name is longer than {limits.max_column_name_chars} characters.",
                hint="Shorten the column names and upload again.",
            )


def _parse_failure(fmt: FileFormat) -> ParseFailure:
    return ParseFailure(
        f"The file could not be read as {fmt.value.upper()}.",
        hint="Check the file opens correctly, has a header row, and every row has the same "
        "number of columns.",
    )


def _detect_encoding(source: Path) -> str:
    """Classify a text file as utf-8, utf-16 (BOM present) or, failing both, cp1252.

    cp1252 is what "CSV (Comma delimited)" from Excel on Windows produces, and it decodes every
    latin-1 file correctly except for a few control-range bytes (replaced, not fatal).
    """
    with source.open("rb") as handle:
        probe = handle.read(_ENCODING_PROBE_BYTES)
    if probe.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    # A multi-byte character may straddle the end of the probe: drop up to 3 trailing bytes.
    for trim in range(4):
        try:
            probe[: len(probe) - trim].decode("utf-8")
        except UnicodeDecodeError:
            continue
        return "utf-8"
    return "cp1252"


def _check_structure(source: Path, fmt: FileFormat, limits: IngestLimits) -> int:
    """Refuse a delimited file whose rows do not line up with its header; returns the header's
    field count.

    A reader that meets a row with too many or too few fields, or a quote that never closes,
    may silently read the whole file as one column (DuckDB's sniffer does). That table would
    be analysed as if it were fine, so the file is checked first, row by row, and refused with
    the line to fix. Messages name lines and counts, never values.
    """
    delimiter = "\t" if fmt is FileFormat.TSV else _delimiter(source)
    if delimiter is None:
        return 1  # one column: nothing to line up
    return _scan_rows(source, delimiter, limits.max_rows)


def _scan_rows(source: Path, delimiter: str, max_rows: int) -> int:
    expected = 0
    first_bad: tuple[int, int] | None = None  # (line, fields)
    bad = 0
    records = 0
    with source.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter, strict=True)
        start = 1
        while True:
            try:
                row = next(reader)
            except StopIteration:
                break
            except csv.Error as exc:
                if first_bad is not None:
                    break  # the earlier line is the one to fix first
                raise ParseFailure(
                    f"A quote opened on line {start:,} is never closed, or a value there has a "
                    "quote in the wrong place.",
                    hint="Fix the quotes on that line (a quote inside a value is written twice: "
                    '""), then upload the file again.',
                ) from exc
            line, start = start, reader.line_num + 1
            if not row:
                continue  # a blank line
            if not expected:
                expected = len(row)
                continue
            records += 1
            if len(row) != expected:
                bad += 1
                first_bad = first_bad or (line, len(row))
            if records > max_rows:
                break  # the row cap is reported by the parse itself
    if first_bad is not None:
        line, fields = first_bad
        others = bad - 1
        more = (
            f" {others:,} other line{'s' if others != 1 else ''} do{'' if others != 1 else 'es'}"
            " not match either."
            if others
            else ""
        )
        raise ParseFailure(
            f"Line {line:,} has {fields} value{'s' if fields != 1 else ''}, but the header has "
            f"{expected} columns.{more}",
            hint="Every row needs the same number of columns. Check that line for an extra or "
            "missing separator or an unclosed quote, then upload the file again.",
        )
    return expected


def _delimiter(source: Path) -> str | None:
    """The separator the header uses, or ``None`` for a one-column file.

    Each candidate splits the first records; the one that gives the header more than one field
    and agrees with the most rows wins (a semicolon file with commas inside quoted names still
    reads as semicolons), then the one with more fields.
    """
    with source.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        sample = handle.read(_ENCODING_PROBE_BYTES)
    lines = sample.splitlines(keepends=True)
    if len(sample) == _ENCODING_PROBE_BYTES and len(lines) > 1:
        lines = lines[:-1]  # the last line may be cut short
    best: tuple[float, int, str] | None = None
    for candidate in _DELIMITERS:
        try:
            rows = [r for r in csv.reader(lines, delimiter=candidate) if r]
        except csv.Error:
            continue
        rows = rows[:_DELIMITER_SAMPLE_RECORDS]
        if not rows or len(rows[0]) < 2:
            continue
        width = len(rows[0])
        body = rows[1:]
        agreement = sum(1 for r in body if len(r) == width) / len(body) if body else 1.0
        score = (agreement, width, candidate)
        if best is None or score[:2] > best[:2]:
            best = score
    return best[2] if best else None


def _transcode_to_utf8(source: Path, dest: Path, encoding: str) -> None:
    with (
        source.open("r", encoding=encoding, errors="replace", newline="") as reader,
        dest.open("w", encoding="utf-8", newline="") as writer,
    ):
        while chunk := reader.read(_HASH_CHUNK):
            writer.write(chunk)


# --------------------------------------------------------------------------- excel


def _xlsx_to_csv(source: Path, dest: Path, limits: IngestLimits) -> None:
    """Stream the first worksheet into a CSV (read-only mode keeps memory flat)."""
    from openpyxl import load_workbook  # noqa: PLC0415  (heavy import, only needed for xlsx)

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt workbooks
        raise ParseFailure(
            "The workbook could not be opened.",
            hint="Open it in Excel and save it again as .xlsx.",
        ) from exc
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ParseFailure("The workbook has no worksheets.")
        with dest.open("w", newline="", encoding="utf-8") as handle:
            _write_sheet(sheet, csv.writer(handle), limits)
    finally:
        workbook.close()


def _write_sheet(sheet: Worksheet, writer: Any, limits: IngestLimits) -> None:
    rows = sheet.iter_rows(values_only=True)
    width = 0
    header_written = False
    data_rows = 0
    for row in rows:
        if not header_written:
            if all(cell is None for cell in row):
                continue  # skip leading blank rows
            width = max(i for i, cell in enumerate(row, start=1) if cell is not None)
            if width > limits.max_columns:
                raise TooManyColumns(
                    f"The first sheet has {width:,} columns; the limit is {limits.max_columns}.",
                    hint="Remove columns you do not need and upload again.",
                )
            writer.writerow(
                [
                    str(cell).strip() if cell is not None else f"column{i}"
                    for i, cell in enumerate(row[:width])
                ]
            )
            header_written = True
            continue
        data_rows += 1
        if data_rows > limits.max_rows:
            raise TooManyRows(
                f"The first sheet has more than {limits.max_rows:,} rows.",
                hint="Upload a smaller file or an aggregated extract.",
            )
        padded = tuple(row[:width]) + (None,) * max(0, width - len(row))
        writer.writerow([_cell_to_text(cell) for cell in padded])
    if not header_written:
        raise ParseFailure("The first sheet is empty.", hint="Check the right sheet is first.")


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


@contextmanager
def _workdir(dest: Path) -> Iterator[Path]:
    """A scratch directory next to the destination, removed on exit."""
    path = Path(tempfile.mkdtemp(prefix="ingest-", dir=dest.parent))
    try:
        yield path
    finally:
        shutil.rmtree(path, ignore_errors=True)
