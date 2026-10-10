"""The Python recipe: a script that rebuilds a run's cleaned data, features and charts.

Cleaned data is always the stored plan applied to the original bytes. The recipe writes that
plan down as code a person can read, audit and run on their own machine: it loads the original
file the way ingest did, pins the recorded column types, applies exactly the accepted cleaning
operations (the same SQL the engine uses), adds the enabled features, and draws the planned
charts with Vega-Lite. Nothing in it was written by a model.
"""

from __future__ import annotations

import json
import pprint
import unicodedata
from collections.abc import Sequence
from typing import Any, Final

from insightfulcell.engine import features as feature_sql
from insightfulcell.engine.cleaning.operations import (
    Decision,
    Operation,
    cleaned_select,
)
from insightfulcell.engine.planning.language import ChartKind, ResolvedPlan, Role
from insightfulcell.engine.sql import ident
from insightfulcell.engine.visualization import data_sql, vega_lite

VEGA: Final = "https://cdn.jsdelivr.net/npm/vega@5.30.0"
VEGA_LITE: Final = "https://cdn.jsdelivr.net/npm/vega-lite@5.21.0"
VEGA_EMBED: Final = "https://cdn.jsdelivr.net/npm/vega-embed@6.26.0"


def _py(value: Any) -> str:
    """A Python literal (``True``, ``None``; not JSON's ``true``, ``null``)."""
    return pprint.pformat(value, width=100, sort_dicts=False)


def _triple(sql: str) -> str:
    """A raw triple-quoted string, readable as SQL. A quoted identifier can contain three
    quotes in a row (a column named ``a"``): then a plain string literal instead."""
    if '"""' in sql:
        return _py(sql)
    return 'r"""\n' + sql + '\n"""'


def _text(value: str) -> str:
    """User text (a file or column name) made safe for a comment or the docstring: no line
    breaks or other control characters (they would end a comment and start code), and no
    quotes or backslashes (they could close or escape the docstring)."""
    out = []
    for ch in value:
        category = unicodedata.category(ch)
        if category[0] == "C" or category in {"Zl", "Zp"}:
            out.append(" ")
        elif ch == '"':
            out.append("'")
        elif ch == "\\":
            out.append("/")
        else:
            out.append(ch)
    return "".join(out)


def render(  # noqa: PLR0913  (keyword-only)
    *,
    dataset_name: str,
    run_id: str,
    source_format: str,
    columns: Sequence[tuple[str, str]],
    read_as_text: bool,
    operations: Sequence[Operation],
    plan: ResolvedPlan,
    cleaning_revision: int,
    cleaning_status: str,
) -> str:
    names = [name for name, _ in columns]
    clean = cleaned_select(names, operations, table="typed")
    accepted = [o for o in operations if o.decision is Decision.ACCEPTED]
    pending = [o for o in operations if o.decision is Decision.PENDING]
    excluded = [c.column for c in plan.dictionary if c.role is Role.EXCLUDED]
    labels = {c.column: c.display_name for c in plan.dictionary}
    kept = [n for n in names if n not in excluded]
    enabled = [f for f in plan.features if f.enabled]
    feature_lines = [
        f"{feature_sql.expression(f.kind, f.columns, f.part, f.bins)} AS {ident(f.name)}"
        for f in enabled
    ]
    analysis = (
        "SELECT "
        + ",\n       ".join([ident(n) for n in kept] + feature_lines or ["*"])
        + "\nFROM cleaned"
    )
    charts = []
    for chart in plan.charts:
        sql = data_sql(chart, table="cleaned")
        if sql is None or chart.kind is ChartKind.MAP:
            continue
        charts.append({"id": chart.id, "sql": sql, "spec": vega_lite(chart, [], labels)})

    steps = (
        "\n".join(
            _text(f"#   {i}. {o.label}: {o.column or 'all rows'}  [{o.id}]")
            for i, o in enumerate(accepted, 1)
        )
        or "#   (none)"
    )
    waiting = (
        "\n".join(_text(f"#   - {o.label}: {o.column or 'all rows'}  [{o.id}]") for o in pending)
        if pending
        else ""
    )
    header = f'''"""InsightfulCell recipe

Dataset:  {_text(dataset_name)}
Run:      {run_id}
Plan:     {plan.digest()[:16]} (cleaning revision {cleaning_revision}, {cleaning_status})

Rebuilds this run's cleaned table, analysis table (features added, excluded columns left out)
and charts from the original file. Generated from the stored plan; nothing here was written by
a language model.

Usage:  python recipe.py ORIGINAL_FILE [OUTPUT_FOLDER]
Needs:  pip install duckdb   (and openpyxl for .xlsx files)
"""'''
    body = f"""
from __future__ import annotations

import csv
import json
import sys
import tempfile
from datetime import date, datetime, time
from pathlib import Path

import duckdb

SOURCE_FORMAT = {json.dumps(source_format)}
# The types recorded when the file was ingested. Values that do not fit become empty.
COLUMNS = {_py([list(c) for c in columns])}
READ_AS_TEXT = {read_as_text!r}  # ingest fell back to reading every column as text
EXCLUDED_FROM_ANALYSIS = {_py(excluded)}

# Cleaning applied, in order (accepted operations only):
{steps}
{"# Waiting for a decision when this recipe was made (not applied):" + chr(10) + waiting if waiting else ""}
CLEAN_SQL = {_triple(clean)}

ANALYSIS_SQL = {_triple(analysis)}

CHARTS = {_py(charts)}


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _xlsx_to_csv(path: Path, dest: Path) -> None:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    with dest.open("w", newline="", encoding="utf-8") as handle:
        writer, width = csv.writer(handle), None
        for row in sheet.iter_rows(values_only=True):
            if width is None:
                if all(cell is None for cell in row):
                    continue
                width = max(i for i, cell in enumerate(row, 1) if cell is not None)
                writer.writerow(
                    [str(c).strip() if c is not None else f"column{{i}}" for i, c in enumerate(row[:width])]
                )
                continue
            padded = tuple(row[:width]) + (None,) * max(0, width - len(row))
            writer.writerow([_cell(c) for c in padded])


def load(con: duckdb.DuckDBPyConnection, path: Path) -> None:
    fmt = SOURCE_FORMAT
    if fmt == "xlsx":
        converted = Path(tempfile.mkdtemp()) / "sheet.csv"
        _xlsx_to_csv(path, converted)
        path, fmt = converted, "csv"
    if fmt == "parquet":
        relation = con.read_parquet(str(path))
    elif fmt == "json":
        relation = con.read_json(str(path))
    else:
        options = {{"sample_size": 100_000, "all_varchar": READ_AS_TEXT}}
        if fmt == "tsv":
            options["delimiter"] = "\\t"
        relation = con.read_csv(str(path), **options)
    relation.create_view("raw")
    typed = ", ".join(
        f'TRY_CAST("{{n.replace(chr(34), chr(34) * 2)}}" AS {{t}}) AS "{{n.replace(chr(34), chr(34) * 2)}}"'
        for n, t in COLUMNS
    )
    con.execute(f"CREATE VIEW typed AS SELECT {{typed}} FROM raw")


def _json_rows(con: duckdb.DuckDBPyConnection, sql: str) -> list[dict]:
    cursor = con.execute(sql)
    names = [d[0] for d in cursor.description]
    return [
        {{k: (v.isoformat() if isinstance(v, (datetime, date)) else v) for k, v in zip(names, row)}}
        for row in cursor.fetchall()
    ]


def charts_html(con: duckdb.DuckDBPyConnection) -> str:
    blocks, scripts = [], []
    for i, chart in enumerate(CHARTS):
        spec = dict(chart["spec"], data={{"values": _json_rows(con, chart["sql"])}})
        blocks.append(f'<section><div id="chart{{i}}" style="width:100%"></div></section>')
        # "<" escaped: a value such as "</script>" in the data must stay data.
        data = json.dumps(spec, default=str).replace("<", "\\\\u003c")
        scripts.append(f"vegaEmbed('#chart{{i}}', {{data}}, {{{{actions: false}}}});")
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>Charts</title>"
        "<script src='{VEGA}'></script><script src='{VEGA_LITE}'></script>"
        "<script src='{VEGA_EMBED}'></script>"
        "<style>body{{font-family:system-ui,sans-serif;max-width:960px;margin:2rem auto;padding:0 1rem}}"
        "section{{margin:0 0 2.5rem}}</style></head><body>"
        + "".join(blocks)
        + "<script>" + "\\n".join(scripts) + "</script></body></html>"
    )


def main(argv: list[str]) -> None:
    if len(argv) < 2:
        sys.exit("usage: python recipe.py ORIGINAL_FILE [OUTPUT_FOLDER]")
    source, out = Path(argv[1]), Path(argv[2] if len(argv) > 2 else "recipe-output")
    out.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect()
    load(con, source)
    con.execute("CREATE TABLE cleaned AS " + CLEAN_SQL)
    con.execute("CREATE TABLE analysis AS " + ANALYSIS_SQL)
    con.execute("COPY cleaned TO ? (FORMAT parquet)", [str(out / "cleaned.parquet")])
    con.execute("COPY analysis TO ? (FORMAT parquet)", [str(out / "analysis.parquet")])
    (out / "charts.html").write_text(charts_html(con), encoding="utf-8")
    rows = con.execute("SELECT count(*) FROM cleaned").fetchone()[0]
    print(f"Wrote {{rows}} cleaned rows, the analysis table and {{len(CHARTS)}} charts to {{out}}/")


if __name__ == "__main__":
    main(sys.argv)
"""
    return header + "\n" + body
